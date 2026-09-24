import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .triage import TriageResult


@dataclass(frozen=True)
class ReportRow:
    source: str
    text: str
    result: TriageResult | None
    error: str = ""


HEADERS = ["Источник", "Категория", "Тема", "Срочность", "Сущности", "Текст", "Ошибка"]
URGENCY_ORDER = {"high": 0, "medium": 1, "low": 2}


def safe_cell(value: str) -> str:
    """Prevent CSV/XLSX formula execution when opening untrusted input in a spreadsheet."""
    if value and value[0] in "=+-@\t\r":
        return "'" + value
    return value


def to_table(rows: list[ReportRow], categories: list[str]) -> list[list[str]]:
    category_order = {category: index for index, category in enumerate(categories)}
    ordered = sorted(
        rows,
        key=lambda row: (
            category_order.get(row.result.category, len(categories))
            if row.result
            else len(categories),
            URGENCY_ORDER.get(row.result.urgency, 3) if row.result else 3,
            row.source,
        ),
    )
    result: list[list[str]] = []
    for row in ordered:
        triage = row.result
        result.append(
            [
                safe_cell(row.source),
                safe_cell(triage.category) if triage else "",
                safe_cell(triage.subject) if triage else "",
                triage.urgency if triage else "",
                safe_cell(", ".join(triage.entities)) if triage else "",
                safe_cell(row.text),
                safe_cell(row.error),
            ]
        )
    return result


def write_report(path: Path, rows: list[ReportRow], categories: list[str]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    data = to_table(rows, categories)
    if path.suffix.lower() == ".csv":
        with path.open("w", newline="", encoding="utf-8-sig") as output:
            writer = csv.writer(output)
            writer.writerow(HEADERS)
            writer.writerows(data)
    elif path.suffix.lower() == ".xlsx":
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Триаж"
        sheet.append(HEADERS)
        for row in data:
            sheet.append(row)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = PatternFill("solid", fgColor="1D3557")
        for column, width in {
            "A": 30,
            "B": 32,
            "C": 36,
            "D": 16,
            "E": 40,
            "F": 80,
            "G": 50,
        }.items():
            sheet.column_dimensions[column].width = width
        workbook.save(path)
    else:
        raise ValueError("Формат отчёта должен быть .csv или .xlsx")